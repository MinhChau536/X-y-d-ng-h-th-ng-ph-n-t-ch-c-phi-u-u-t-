import numpy as np, pandas as pd, pytest
import fetch_financials as f
from ratios import get_ratios

Y = [2021, 2022, 2023, 2024, 2025]; B = 1e9
COLS = ["ticker", "year", "item_code", "item_name", "value"]

def long(t, st, items):
    if not items: return pd.DataFrame(columns=COLS)
    return pd.DataFrame([dict(ticker=t, year=y, item_code=f"{st[:2]}{i}", item_name=n, value=v * B)
                         for i, (n, vs) in enumerate(items.items()) for y, v in zip(Y, vs)])

CORP = {
 "balance_sheet": {"TỔNG TÀI SẢN": [900, 1000, 1100, 1200, 1300], "TÀI SẢN NGẮN HẠN": [500, 550, 600, 650, 700],
   "Tiền và tương đương tiền": [100, 110, 120, 130, 140], "Hàng tồn kho, ròng": [200, 220, 240, 260, 280],
   "Các khoản phải thu": [100, 110, 120, 130, 140], "NỢ PHẢI TRẢ": [400, 450, 500, 550, 600],
   "Nợ ngắn hạn": [300, 330, 360, 390, 420], "Vay ngắn hạn": [100] * 5, "Vay dài hạn": [50] * 5,
   "VỐN CHỦ SỞ HỮU": [500, 550, 600, 650, 700]},
 "income_statement": {"Doanh số thuần": [800, 880, 968, 1000, 1100], "Giá vốn hàng bán": [-600, -660, -726, -750, -800],
   "Lãi gộp": [200, 220, 242, 250, 300], "Trong đó: Chi phí lãi vay": [-10] * 5,
   "Lãi/(lỗ) ròng trước thuế": [90, 100, 110, 120, 130], "Lãi/(lỗ) thuần sau thuế": [70, 80, 90, 100, 110],
   "EBITDA": [150, 160, 170, 180, 190]},
 "cash_flow": {"Lưu chuyển tiền thuần từ các hoạt động sản xuất kinh doanh": [80, 90, 100, 110, 120],
   "Tiền mua tài sản cố định và các tài sản dài hạn khác": [-50] * 5}}
BANK = {
 "balance_sheet": {"TỔNG TÀI SẢN": [1000, 1100, 1200, 1300, 1400], "Tổng nợ phải trả": [920, 1010, 1100, 1190, 1280],
   "Vốn chủ sở hữu": [80, 90, 100, 110, 120], "Tiền gửi của khách hàng": [700, 770, 840, 910, 980]},
 "income_statement": {"Thu nhập lãi thuần": [40, 44, 48, 52, 56], "Tổng thu nhập hoạt động": [50, 55, 60, 65, 70],
   "Chi phí hoạt động": [-20, -21, -22, -23, -24], "Chi phí dự phòng rủi ro tín dụng": [-5] * 5,
   "Tổng lợi nhuận trước thuế": [25, 29, 33, 37, 41], "Lợi nhuận sau thuế": [20, 23, 26, 29, 33]},
 "cash_flow": {}}

@pytest.fixture(autouse=True)
def fake(monkeypatch):
    stores = {st: pd.concat([long("TST", st, CORP[st]), long("BNK", st, BANK[st])], ignore_index=True) for st in f.STATEMENTS}
    empty = pd.DataFrame(columns=COLS)
    monkeypatch.setattr(f, "_load_raw", lambda ex, s: stores[s] if ex == "HSX" else empty)
    f._load.cache_clear(); yield; f._load.cache_clear()

def test_corporate_ratios():
    r = get_ratios(f.get_financials("TST"))
    x = r["ratios"]
    assert r["sector"] == "corporate"
    assert x.loc[2025, "ROE (BQ)"] == pytest.approx(110 / 675)
    assert x.loc[2025, "ROE (cuối kỳ)"] == pytest.approx(110 / 700)
    assert x.loc[2025, "Tăng trưởng LNST"] == pytest.approx(0.10)
    assert x.loc[2025, "Thanh toán hiện hành"] == pytest.approx(700 / 420)
    assert x.loc[2025, "Nợ/Vốn CSH"] == pytest.approx(600 / 700)
    assert x.loc[2025, "CFO/LNST"] == pytest.approx(120 / 110)
    assert x.loc[2025, "FCF/Doanh thu"] == pytest.approx(70 / 1100)
    assert x.loc[2025, "Nợ vay ròng/EBITDA"] == pytest.approx(10 / 190)
    assert np.isnan(x.loc[2021, "ROE (BQ)"])            # năm đầu chưa có số dư bình quân

def test_bank_ratios():
    r = get_ratios(f.get_financials("BNK"))
    assert r["sector"] == "bank"
    assert r["ratios"].loc[2025, "CIR"] == pytest.approx(24 / 70)
    assert "Thanh toán hiện hành" not in r["ratios"].columns

def test_missing_cashflow_is_nan_and_n_years():
    d = f.get_financials("BNK")
    r = get_ratios(d, n_years=3)
    assert list(r["ratios"].index) == [2023, 2024, 2025]
    t = get_ratios(f.get_financials("TST"))
    assert (t["mapping"]["Chỉ tiêu trong dataset"] == "— (không có dữ liệu)").sum() >= 1    # vd. ebit không có trong dữ liệu

def test_cli_writes_excel_with_ratios(tmp_path, capsys):
    d = f.get_financials("TST")
    from ratios import PERCENT
    rr = get_ratios(d)
    p = f.save_excel(d, str(tmp_path), {"Chỉ số tài chính": (rr["ratios"].T, PERCENT)})
    assert "Chỉ số tài chính" in pd.ExcelFile(p).sheet_names
    import openpyxl
    ws = openpyxl.load_workbook(p)["Chỉ số tài chính"]
    fmts = {ws.cell(r, 1).value: ws.cell(r, 3).number_format for r in range(2, ws.max_row + 1)}
    assert fmts["ROE (BQ)"] == "0.0%" and fmts["Nợ/Vốn CSH"] == "#,##0.00"
